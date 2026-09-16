"""CarryOn™ — Partner roster import (spreadsheet → mapping → reconciliation).

Partners keep a *living* spreadsheet of their clients: columns move, headers
get renamed, the same people appear on every upload. This module turns any
such sheet into a deterministic plan against the partner's existing roster:

  identity  — the client's EMAIL is the key (it is where the claim invitation
              goes). Never AI. Same person → same answer on every upload.
  layout    — column mapping: remembered per partner (header fingerprint) →
              deterministic detection → optional xAI suggestion on masked
              samples → always editable by the partner.
  safety    — a spreadsheet can only ADD clients or rename unclaimed
              portals. It never deletes, deactivates or touches claimed
              accounts ("roster wins").
"""

import asyncio
import csv
import hashlib
import io
import json
import re
import time
from datetime import datetime, timezone

from config import XAI_MODEL_LIGHT, db, logger, xai_client
from services.email import EMAIL_REGEX, is_valid_email

MAX_ROWS = 2000
MAX_FILE_BYTES = 5 * 1024 * 1024
SAMPLE_ROWS = 3
FIELDS = ("email", "first_name", "last_name", "full_name")

_HEADER_HINTS = {
    "email": (
        "email",
        "emailaddress",
        "clientemail",
        "primaryemail",
        "mail",
        "emailprimary",
        "memberemail",
        "contactemail",
    ),
    "first_name": (
        "firstname",
        "first",
        "fname",
        "givenname",
        "given",
        "clientfirstname",
        "firstnm",
        "memberfirstname",
    ),
    "last_name": ("lastname", "last", "lname", "surname", "familyname", "clientlastname", "lastnm", "memberlastname"),
    "full_name": (
        "name",
        "fullname",
        "clientname",
        "client",
        "contact",
        "contactname",
        "customer",
        "customername",
        "member",
        "membername",
        "participant",
        "participantname",
        "employee",
        "employeename",
    ),
}
_SECONDARY_WORDS = ("alt", "alternate", "secondary", "spouse", "partner", "2", "other", "backup", "assistant")

REASON_TEXT = {
    "new": "Will be added",
    "name_update": "Name will be updated (portal not claimed yet)",
    "already_client": "Already a client",
    "existing_account": "A CarryOn account already exists for this email (not in your roster)",
    "duplicate_in_file": "Appears earlier in this file",
    "needs_email": "No email address",
    "invalid_email": "Email address doesn't look valid",
    "needs_name": "First and last name are both required",
    "no_seat": "No seat left in your partnership",
}


def normalize_header(h: str) -> str:
    return re.sub(r"[^a-z0-9]", "", (h or "").lower())


def fingerprint(headers: list[str]) -> str:
    """Order-independent layout id — reordering columns keeps the remembered mapping."""
    key = "|".join(sorted(normalize_header(h) for h in headers if normalize_header(h)))
    return hashlib.sha1(key.encode()).hexdigest()[:16]


def _clean(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def parse_spreadsheet(filename: str, content: bytes) -> tuple[list[str], list[list[str]], str]:
    """→ (headers, rows, sheet_name). Title/banner rows are skipped; CSV or XLSX."""
    if len(content) > MAX_FILE_BYTES:
        raise ValueError("That file is larger than 5 MB. Remove extra sheets or images and try again.")
    name = (filename or "").lower()
    sheet_name = ""
    if name.endswith((".xlsx", ".xlsm")):
        table, sheet_name = _read_xlsx(content)
    elif name.endswith((".csv", ".txt")):
        table = _read_csv(content)
    elif name.endswith(".xls"):
        raise ValueError("Legacy .xls files aren't supported — save the sheet as .xlsx or .csv and try again.")
    else:
        raise ValueError("Upload a .csv or .xlsx file.")

    table = [[_clean(c) for c in row] for row in table]
    table = [row for row in table if any(row)]
    if len(table) < 2:
        raise ValueError(
            "We couldn't find any people in that file. It needs a header row (e.g. First Name, Last Name, Email) followed by one row per person."
        )
    start = _header_row_index(table)
    headers = table[start]
    width = max(len(r) for r in table[start:])
    headers = headers + [f"Column {i + 1}" for i in range(len(headers), width)]
    headers = [h or f"Column {i + 1}" for i, h in enumerate(headers)]
    rows = [r + [""] * (width - len(r)) for r in table[start + 1 :]]
    if not rows:
        raise ValueError("That file only has a header row — add one row per person and try again.")
    if len(rows) > MAX_ROWS:
        raise ValueError(
            f"Too many rows — the limit is {MAX_ROWS:,} per upload. Split the sheet and import it in parts."
        )
    return headers, rows, sheet_name


def _header_row_index(table: list[list[str]]) -> int:
    """Skip title/banner rows: the header is the first row (in the top 10) with
    2+ filled cells and no email-looking value."""
    for i, row in enumerate(table[:10]):
        filled = [c for c in row if c]
        if len(filled) >= 2 and not any(EMAIL_REGEX.match(c) for c in filled):
            return i
    return 0


def _read_csv(content: bytes) -> list[list]:
    text = content.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return list(csv.reader(io.StringIO(text), dialect))


def _read_xlsx(content: bytes) -> tuple[list[list], str]:
    """Reads the sheet with the most filled rows (people usually live on one tab)."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001
        raise ValueError("That doesn't look like a valid Excel file. Save it as .xlsx and try again.")
    best: list[list] = []
    best_name = ""
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True, max_row=MAX_ROWS + 50)]
        rows = [r for r in rows if any(_clean(c) for c in r)]
        if len(rows) > len(best):
            best, best_name = rows, ws.title
    wb.close()
    return best, best_name


# ─── Mapping ─────────────────────────────────────────────────────────────


def _email_ratio(rows: list[list[str]], idx: int) -> float:
    vals = [r[idx] for r in rows if idx < len(r) and r[idx]]
    if not vals:
        return 0.0
    return sum(1 for v in vals if EMAIL_REGEX.match(v)) / len(vals)


def _is_secondary(h: str) -> bool:
    n = normalize_header(h)
    return any(w in n for w in _SECONDARY_WORDS)


def detect_mapping(headers: list[str], rows: list[list[str]]) -> dict:
    """Deterministic header → field detection. Values are header NAMES (or None)."""
    mapping: dict = {f: None for f in FIELDS}
    norm = [normalize_header(h) for h in headers]

    email_candidates = [i for i in range(len(headers)) if _email_ratio(rows, i) >= 0.6]
    if email_candidates:
        primary = [i for i in email_candidates if not _is_secondary(headers[i])]
        pool = primary or email_candidates
        pool.sort(key=lambda i: (0 if "email" in norm[i] else 1, i))
        mapping["email"] = headers[pool[0]]
    else:
        for i, n in enumerate(norm):
            if n in _HEADER_HINTS["email"] and not _is_secondary(headers[i]):
                mapping["email"] = headers[i]
                break

    for field in ("first_name", "last_name", "full_name"):
        for i, n in enumerate(norm):
            if n in _HEADER_HINTS[field] and headers[i] != mapping["email"] and not _is_secondary(headers[i]):
                mapping[field] = headers[i]
                break
    if mapping["first_name"] and mapping["last_name"]:
        mapping["full_name"] = None
    return mapping


def mapping_complete(mapping: dict) -> bool:
    return bool(mapping.get("email")) and bool(
        (mapping.get("first_name") and mapping.get("last_name")) or mapping.get("full_name")
    )


def split_full_name(value: str) -> tuple[str, str]:
    v = re.sub(r"\s+", " ", (value or "").strip())
    if not v:
        return "", ""
    if "," in v:  # "Dawson, Jane"
        last, _, first = v.partition(",")
        return first.strip(), last.strip()
    parts = v.split(" ")
    if len(parts) == 1:
        return parts[0], ""
    return " ".join(parts[:-1]), parts[-1]


def apply_mapping(headers: list[str], rows: list[list[str]], mapping: dict) -> list[dict]:
    idx = {f: headers.index(mapping[f]) for f in FIELDS if mapping.get(f) in headers}

    def cell(row, f):
        i = idx.get(f)
        return row[i] if i is not None and i < len(row) else ""

    people = []
    for n, row in enumerate(rows, start=2):  # row 1 is the header
        first, last = cell(row, "first_name"), cell(row, "last_name")
        if not (first and last) and "full_name" in idx:
            f2, l2 = split_full_name(cell(row, "full_name"))
            first, last = first or f2, last or l2
        people.append({"row": n, "first_name": first[:60], "last_name": last[:60], "email": cell(row, "email")})
    return people


def mask(value: str) -> str:
    v = (value or "").strip()
    if not v:
        return ""
    if EMAIL_REGEX.match(v):
        local, _, domain = v.partition("@")
        return f"{local[0]}***@{domain}"
    if re.fullmatch(r"[\d\s()+\-.]+", v):
        return re.sub(r"\d", "#", v)
    return v[0] + "***"


async def ai_suggest_mapping(headers: list[str], rows: list[list[str]], *, actor_id: str | None) -> dict | None:
    """Ask the light xAI model to map headers → fields using MASKED samples.
    Returns a validated mapping or None (the caller falls back to manual)."""
    if xai_client is None:
        return None
    samples = [[mask(c) for c in r] for r in rows[:SAMPLE_ROWS]]
    prompt = (
        "You map spreadsheet columns for importing a list of clients. Fields:\n"
        "- email: the client's PRIMARY email (not spouse/alternate/assistant)\n"
        "- first_name, last_name: separate name columns, if present\n"
        "- full_name: a single column holding the whole name, if there are no separate name columns\n"
        f"Column headers: {json.dumps(headers)}\n"
        f"Masked sample rows: {json.dumps(samples)}\n"
        'Reply with ONLY a JSON object like {"email": "<header or null>", "first_name": ..., "last_name": ..., "full_name": ...} '
        "using the exact header text. Use null when no column fits."
    )
    t0 = time.time()
    try:
        resp = await asyncio.to_thread(
            xai_client.chat.completions.create,
            model=XAI_MODEL_LIGHT,
            messages=[{"role": "user", "content": prompt}],
            temperature=0,
            max_tokens=300,
        )
        try:
            from services.llm_cost_ledger import record_xai_response

            await record_xai_response(
                resp, endpoint="roster.map_columns", model=XAI_MODEL_LIGHT, user_id=actor_id, started_at=t0
            )
        except Exception:  # noqa: BLE001
            pass
        text = (resp.choices[0].message.content or "").strip()
        if text.startswith("```"):
            text = text.split("```", 2)[1]
            text = text[4:] if text.startswith("json") else text
        parsed = json.loads(text)
    except Exception as e:  # noqa: BLE001
        logger.warning("Roster AI mapping failed: %s", e)
        return None
    mapping = {f: (parsed.get(f) if parsed.get(f) in headers else None) for f in FIELDS}
    # Sanity-check the suggestion against the data itself: the email column must
    # actually hold emails, name columns must not.
    if mapping["email"] and _email_ratio(rows, headers.index(mapping["email"])) < 0.5:
        mapping["email"] = None
    for f in ("first_name", "last_name", "full_name"):
        if mapping[f] and _email_ratio(rows, headers.index(mapping[f])) > 0.3:
            mapping[f] = None
    if mapping["first_name"] and mapping["last_name"]:
        mapping["full_name"] = None
    return mapping if mapping["email"] else None


async def remembered_mapping(partner: dict, headers: list[str]) -> dict | None:
    saved = (partner.get("roster_layouts") or {}).get(fingerprint(headers))
    if not saved:
        return None
    mapping = {f: saved.get("mapping", {}).get(f) for f in FIELDS}
    return mapping if all(v is None or v in headers for v in mapping.values()) and mapping_complete(mapping) else None


async def remember_mapping(partner_id: str, headers: list[str], mapping: dict) -> None:
    await db.b2b_partners.update_one(
        {"id": partner_id},
        {
            "$set": {
                f"roster_layouts.{fingerprint(headers)}": {
                    "mapping": {f: mapping.get(f) for f in FIELDS},
                    "headers": headers[:60],
                    "saved_at": datetime.now(timezone.utc).isoformat(),
                }
            }
        },
    )


# ─── Reconciliation ──────────────────────────────────────────────────────


async def reconcile(partner: dict, people: list[dict]) -> dict:
    """Plan for this upload against the partner's roster. Pure read; nothing is written."""
    roster = await db.users.find(
        {"partner_id": partner["id"]},
        {
            "_id": 0,
            "id": 1,
            "email_lower": 1,
            "email": 1,
            "name": 1,
            "first_name": 1,
            "last_name": 1,
            "account_status": 1,
        },
    ).to_list(5000)
    by_email = {(u.get("email_lower") or (u.get("email") or "").lower()): u for u in roster}

    sheet_emails = [p["email"].strip().lower() for p in people if EMAIL_REGEX.match(p["email"].strip())]
    foreign = set()
    if sheet_emails:
        others = await db.users.find(
            {
                "$or": [
                    {"email_lower": {"$in": sheet_emails}},
                    {"email": {"$in": sheet_emails}},
                    {"username_lower": {"$in": sheet_emails}},
                ]
            },
            {"_id": 0, "id": 1, "email_lower": 1, "email": 1, "partner_id": 1},
        ).to_list(5000)
        foreign = {
            (u.get("email_lower") or (u.get("email") or "").lower())
            for u in others
            if u.get("partner_id") != partner["id"]
        }

    max_uses = int(partner.get("max_uses", 0) or 0)
    seats_left = max(0, max_uses - int(partner.get("times_used", 0) or 0)) if max_uses > 0 else None

    seen: set = set()
    rows = []
    for p in people:
        email_raw = p["email"].strip()
        email = email_raw.lower()
        row = {**p, "email": email_raw, "action": "skip", "reason": "", "client_id": None, "current_name": None}
        if not email:
            row["reason"] = "needs_email"
        elif not EMAIL_REGEX.match(email) or not is_valid_email(email):
            row["reason"] = "invalid_email"
        elif email in seen:
            row["reason"] = "duplicate_in_file"
        elif email in by_email:
            u = by_email[email]
            row["client_id"], row["current_name"] = u["id"], u.get("name", "")
            new_name = f"{p['first_name']} {p['last_name']}".strip()
            if (
                u.get("account_status") == "pending_claim"
                and p["first_name"]
                and p["last_name"]
                and new_name != (u.get("name") or "").strip()
            ):
                row["action"], row["reason"] = "update_name", "name_update"
            else:
                row["reason"] = "already_client"
        elif email in foreign:
            row["reason"] = "existing_account"
        elif not (p["first_name"] and p["last_name"]):
            row["reason"] = "needs_name"
        elif seats_left is not None and seats_left <= 0:
            row["reason"] = "no_seat"
        else:
            row["action"], row["reason"] = "add", "new"
            if seats_left is not None:
                seats_left -= 1
        if email:
            seen.add(email)
        row["detail"] = REASON_TEXT.get(row["reason"], "")
        rows.append(row)

    in_file = {r["email"].lower() for r in rows if r["email"]}
    not_in_upload = [
        {
            "name": u.get("name", ""),
            "email": u.get("email", ""),
            "status": "pending_claim" if u.get("account_status") == "pending_claim" else "active",
        }
        for e, u in by_email.items()
        if e and e not in in_file
    ]
    counts: dict = {}
    for r in rows:
        counts[r["reason"]] = counts.get(r["reason"], 0) + 1
    return {
        "rows": rows,
        "summary": {
            "total_rows": len(rows),
            "add": counts.get("new", 0),
            "update_name": counts.get("name_update", 0),
            "already_client": counts.get("already_client", 0),
            "existing_account": counts.get("existing_account", 0),
            "duplicate_in_file": counts.get("duplicate_in_file", 0),
            "needs_email": counts.get("needs_email", 0) + counts.get("invalid_email", 0),
            "needs_name": counts.get("needs_name", 0),
            "no_seat": counts.get("no_seat", 0),
            "not_in_upload": len(not_in_upload),
            "seats_remaining": seats_left,
        },
        "not_in_upload": sorted(not_in_upload, key=lambda x: x["name"])[:200],
    }
